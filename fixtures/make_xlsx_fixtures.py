#!/usr/bin/env python3
"""Regenerate the synthetic .xlsx fixtures. Every serial here is made up -
never put a real customer's license file in this (public) repo.

sample-spreadsheet.xlsx      the shapes people actually send: Wizard lines
                             pasted into cells beside notes columns, a table
                             with License / Qty / Serial headers, bare keys
                             next to names with no header at all.
sample-spreadsheet-raw.xlsx  hand-written SpreadsheetML that openpyxl would
                             never produce - namespace prefixes, inline and
                             rich strings, rows/cells with no r= reference,
                             stored (undeflated) parts, an absolute rel target.
"""
import zipfile
from pathlib import Path

from openpyxl import Workbook

here = Path(__file__).resolve().parent


def sample():
    wb = Workbook()
    ws = wb.active
    ws.title = "EU GC + LC"
    ws.append(["Licenses for site to be refreshed and activated"])
    ws.append([])
    ws.append(["Agilent GC Licenses + Serial Numbers", "Counts", "Source DB", "Target DB"])
    first = ws.max_row + 1
    ws.append(["[Empower 3 Agilent GC Ctrl License Pack] System Licenses: 1 Serial No: K3TAQ1001M", 1, "SRCDB001", "TGTDB002"])
    ws.append(["[Empower 3 Agilent GC Ctrl License Pack] System Licenses: 1 Serial No: K3TAQ1002M", 1, "SRCDB001", "TGTDB002"])
    # leading bracket lost in the copy/paste
    ws.append(["Empower 3 Agilent GC Ctrl License Pack] System Licenses: 20 Serial No: K3TAQ1003M", 20, "SRCDB001", "TGTDB002"])
    # the typed Counts column disagrees with the line - the line wins
    ws.append(["[Empower 3 Agilent GC Ctrl License Pack] System Licenses: 5 Serial No: K3TAQ1004M", 1, "SRCDB001", "TGTDB002"])
    # pack suffix + a duplicate of an earlier key
    ws.append(["[Empower 3 Agilent GC Ctrl License Pack] System Licenses: 1 Serial No: K3TAQ1005M-001", 1, "SRCDB001", "TGTDB002"])
    ws.append(["[Empower 3 Agilent GC Ctrl License Pack] System Licenses: 1 Serial No: K3TAQ1001M", 1, "SRCDB009", "TGTDB002"])
    ws.append(["Total Counts:", f"=SUM(B{first}:B{ws.max_row})"])
    ws.append([])
    ws.append(["Shimadzu LC Licenses", "Counts"])
    # several lines in one cell, one of them Checksum-style
    ws.append(["[Empower 3 Shimadzu LC Ctrl License Pack] Instrument control licenses: 2 Serial No: K4SHM2001L\n"
               "Option Properly Installed - Shimadzu LC License(s)   Serial Number - K4SHM2002L", 3])
    # no brackets at all
    ws.append(["Empower 3 Thermo LC Ctrl License Pack System Licenses: 1 Serial No: K5THM3001L"])
    # a key with nothing to read - reported, not dropped
    ws.append(["[Empower 3 Hitachi LC Ctrl License Pack] System Licenses: 1 Serial No:", 1])

    t = wb.create_sheet("Options")
    t.append(["Customer options list"])
    t.append(["License", "Qty", "Serial Number", "Notes"])
    t.append(["Empower 3 Base License Pack", 1, "R7BAS0001X", "server"])
    t.append(["Empower 3 Named User License Pack", 5, "R7BAS0001X-001", "bundled with base"])
    t.append(["Empower 3 System Suitability", None, "Q1SSU2345A", None])
    t.append(["Empower 3 SystemsQT License", 1, "Q1SQT2345B", None])
    t.append(["Empower 3 Agilent GC Ctrl License Pack", 1, "K3TAQ1002M", "also on sheet 1"])
    t.append([None, 1, "Q1NON9999C", "see PO 4471"])

    n = wb.create_sheet("No header")
    n.append(["Empower 3 GPC/SEC Option", "Z9GPC1234X"])
    n.append(["Empower 3 Dissolution Option", "Z9DIS1234Y", "installed 2024"])
    n.append(["Instrument S/N", "AB12CD34EF"])
    n.append(["Lab B notes", "nothing to see here"])
    wb.save(here / "sample-spreadsheet.xlsx")


RAW_WORKBOOK = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<x:workbook xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<x:sheets><x:sheet name="Keys &amp; Packs" sheetId="1" r:id="rId1"/><x:sheet name='Second' sheetId="2" r:id="rId2"/></x:sheets></x:workbook>"""
RAW_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/a.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/../worksheets/b.xml"/>
<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="strings.xml"/>
</Relationships>"""
RAW_STRINGS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<x:sst xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="4" uniqueCount="4">
<x:si><x:r><x:rPr><x:b/></x:rPr><x:t>[Empower 3 </x:t></x:r><x:r><x:t xml:space="preserve">Agilent LC Ctrl License Pack] </x:t></x:r><x:r><x:t>System licenses: 3 Serial No: M1AGL4001Q</x:t></x:r><x:rPh sb="0" eb="1"><x:t>IGNORED</x:t></x:rPh></x:si>
<x:si><x:t>License key</x:t></x:si>
<x:si/>
<x:si><x:t>Empower 3 Base License &amp; Named User Pack</x:t></x:si>
</x:sst>"""
RAW_SHEET_A = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<x:worksheet xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><x:cols><x:col min="1" max="2" width="40"/></x:cols>
<x:sheetData>
<x:row><x:c t="s"><x:v>0</x:v></x:c><x:c><x:v>3</x:v></x:c></x:row>
<x:row><x:c t="inlineStr"><x:is><x:t>[Empower 3 PerkinElmer GC Ctrl License Pack] System licenses: 1 Serial No: M2PKE5001Q_x000D_</x:t></x:is></x:c></x:row>
<x:row r="7"><x:c r="B7" t="s"><x:v>1</x:v></x:c><x:c r="C7" t="inlineStr"><x:is><x:t>Product</x:t></x:is></x:c><x:c r="D7" t="str"><x:f>"Q"&amp;"ty"</x:f><x:v>Qty</x:v></x:c></x:row>
<x:row r="8"><x:c r="B8" t="inlineStr"><x:is><x:t>M3BAS6001Q</x:t></x:is></x:c><x:c r="C8" t="s"><x:v>3</x:v></x:c><x:c r="D8"><x:v>2</x:v></x:c></x:row>
<x:row r="9"><x:c r="B9" t="inlineStr"><x:is><x:t>M3BAS6002Q</x:t></x:is></x:c><x:c r="C9" t="s"><x:v>2</x:v></x:c><x:c r="D9" t="b"><x:v>1</x:v></x:c></x:row>
<x:row r="10"><x:c r="B10" t="e"><x:v>#REF!</x:v></x:c><x:c r="C10"><x:f>SUM(D8:D9)</x:f></x:c></x:row>
</x:sheetData></x:worksheet>"""
RAW_SHEET_B = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>
<row r="1"><c r="A1" t="inlineStr"><is><t>Option Properly Installed - 1 Dissolution License(s)          Serial Number - M4DIS7001Q-002</t></is></c></row>
<row r="2"/>
<row r="3"><c r="AA3" t="inlineStr"><is><t>[Empower 3 GPC Option] Serial No: M4GPC7002Q</t></is></c></row>
</sheetData></worksheet>"""
RAW_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/></Types>"""


def raw():
    with zipfile.ZipFile(here / "sample-spreadsheet-raw.xlsx", "w") as z:
        z.writestr("[Content_Types].xml", RAW_TYPES, zipfile.ZIP_DEFLATED)
        z.writestr("xl/workbook.xml", RAW_WORKBOOK, zipfile.ZIP_STORED)
        z.writestr("xl/_rels/workbook.xml.rels", RAW_RELS, zipfile.ZIP_DEFLATED)
        z.writestr("xl/strings.xml", RAW_STRINGS, zipfile.ZIP_DEFLATED)
        z.writestr("xl/worksheets/a.xml", RAW_SHEET_A, zipfile.ZIP_DEFLATED)
        z.writestr("xl/worksheets/b.xml", RAW_SHEET_B, zipfile.ZIP_STORED)


if __name__ == "__main__":
    sample()
    raw()
    print("wrote", here / "sample-spreadsheet.xlsx", here / "sample-spreadsheet-raw.xlsx")
